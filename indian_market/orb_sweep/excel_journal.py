"""
Comprehensive Excel trade journal for ORB Sweep + Straddle system.

Tracks every trade with full context — signal reason, premiums, greeks,
market state, R-multiples — and auto-generates daily/weekly/monthly
summaries plus a dashboard sheet.

Uses openpyxl for .xlsx output.
"""

import logging
import math
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Dict, List, Optional

import numpy as np

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import (
        Alignment, Border, Font, PatternFill, Side, numbers,
    )
    from openpyxl.utils import get_column_letter
except ImportError:
    raise ImportError(
        "openpyxl is required for the Excel journal. "
        "Install it: pip install openpyxl"
    )

from orb_sweep.config import TradingConfig

logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# Style constants
# ---------------------------------------------------------------------------

_HEADER_FONT = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="2F5496")
_HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
_THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
_GREEN_FILL = PatternFill("solid", fgColor="C6EFCE")
_RED_FILL = PatternFill("solid", fgColor="FFC7CE")
_LIGHT_BLUE_FILL = PatternFill("solid", fgColor="D6E4F0")
_BOLD_FONT = Font(name="Calibri", bold=True, size=11)
_METRIC_FONT = Font(name="Calibri", bold=True, size=12)
_TITLE_FONT = Font(name="Calibri", bold=True, size=14, color="2F5496")

_RUPEE_FMT = '#,##0'
_RUPEE_FMT_2 = '#,##0.00'
_PCT_FMT = '0.00%'
_R_FMT = '0.00"R"'


# ---------------------------------------------------------------------------
# Column definitions for the Trades sheet
# ---------------------------------------------------------------------------

TRADE_COLUMNS = [
    ("Trade #", 8),
    ("Date", 12),
    ("Entry Time", 18),
    ("Exit Time", 18),
    ("Strategy", 14),
    ("Signal Reason", 22),
    ("Direction", 10),
    ("Strike", 10),
    ("Spot Entry", 12),
    ("Spot Exit", 12),
    ("Entry Premium", 14),
    ("Exit Premium", 14),
    ("SL Amount", 12),
    ("TP Amount", 12),
    ("India VIX", 10),
    ("BS Delta", 10),
    ("Expiry Day?", 11),
    ("ORB High", 12),
    ("ORB Low", 12),
    ("VWAP Entry", 12),
    ("Quantity", 10),
    ("Lots", 6),
    ("Gross P&L", 12),
    ("Costs", 10),
    ("Net P&L", 12),
    ("R-Value", 10),
    ("R-Multiple", 10),
    ("Exit Type", 10),
    ("Running Equity", 14),
]


# ---------------------------------------------------------------------------
# ExcelJournal class
# ---------------------------------------------------------------------------

class ExcelJournal:
    """
    Append-only Excel trade journal.

    Usage:
        journal = ExcelJournal(config)
        journal.log_trade(trade_data)
        journal.rebuild_summaries()  # call at EOD
    """

    SHEET_TRADES = "Trades"
    SHEET_DAILY = "Daily Summary"
    SHEET_WEEKLY = "Weekly Summary"
    SHEET_MONTHLY = "Monthly Summary"
    SHEET_DASHBOARD = "Dashboard"

    def __init__(self, config: TradingConfig):
        self.cfg = config
        data_dir = Path(config.db_path).parent
        data_dir.mkdir(parents=True, exist_ok=True)
        self.path = data_dir / "trade_journal.xlsx"
        self._ensure_workbook()

    # ------------------------------------------------------------------
    # Workbook lifecycle
    # ------------------------------------------------------------------

    def _ensure_workbook(self):
        """Create workbook with all sheets if it does not exist."""
        if self.path.exists():
            try:
                wb = load_workbook(str(self.path))
                # Ensure all sheets exist
                for name in [
                    self.SHEET_TRADES, self.SHEET_DAILY,
                    self.SHEET_WEEKLY, self.SHEET_MONTHLY,
                    self.SHEET_DASHBOARD,
                ]:
                    if name not in wb.sheetnames:
                        wb.create_sheet(name)
                wb.save(str(self.path))
                wb.close()
                return
            except Exception as e:
                logger.warning(f"Existing workbook corrupt, recreating: {e}")

        wb = Workbook()
        # Rename default sheet to Trades
        ws = wb.active
        ws.title = self.SHEET_TRADES
        self._setup_trades_sheet(ws)

        for name in [
            self.SHEET_DAILY, self.SHEET_WEEKLY,
            self.SHEET_MONTHLY, self.SHEET_DASHBOARD,
        ]:
            wb.create_sheet(name)

        self._setup_daily_sheet(wb[self.SHEET_DAILY])
        self._setup_weekly_sheet(wb[self.SHEET_WEEKLY])
        self._setup_monthly_sheet(wb[self.SHEET_MONTHLY])
        self._setup_dashboard_sheet(wb[self.SHEET_DASHBOARD])

        wb.save(str(self.path))
        wb.close()
        logger.info(f"Created Excel journal: {self.path}")

    # ------------------------------------------------------------------
    # Sheet setup helpers
    # ------------------------------------------------------------------

    def _apply_header_row(self, ws, columns: list):
        """Write header row and set column widths."""
        for col_idx, (name, width) in enumerate(columns, 1):
            cell = ws.cell(row=1, column=col_idx, value=name)
            cell.font = _HEADER_FONT
            cell.fill = _HEADER_FILL
            cell.alignment = _HEADER_ALIGN
            cell.border = _THIN_BORDER
            ws.column_dimensions[get_column_letter(col_idx)].width = width
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    def _setup_trades_sheet(self, ws):
        self._apply_header_row(ws, TRADE_COLUMNS)

    def _setup_daily_sheet(self, ws):
        cols = [
            ("Date", 12), ("Trades", 8), ("Wins", 8), ("Losses", 8),
            ("Win Rate", 10), ("Gross P&L", 12), ("Net P&L", 12),
            ("Max Intraday DD", 14), ("Equity EOD", 14),
        ]
        self._apply_header_row(ws, cols)

    def _setup_weekly_sheet(self, ws):
        cols = [
            ("Week Start", 12), ("Week End", 12), ("Trades", 8),
            ("Wins", 8), ("Losses", 8), ("Win Rate", 10),
            ("Net P&L", 12), ("Avg R-Multiple", 12), ("Equity EOW", 14),
        ]
        self._apply_header_row(ws, cols)

    def _setup_monthly_sheet(self, ws):
        cols = [
            ("Month", 12), ("Trades", 8), ("Wins", 8), ("Losses", 8),
            ("Win Rate", 10), ("Net P&L", 12), ("Profit Factor", 12),
            ("Max DD", 12), ("Equity EOM", 14),
        ]
        self._apply_header_row(ws, cols)

    def _setup_dashboard_sheet(self, ws):
        ws.column_dimensions["A"].width = 25
        ws.column_dimensions["B"].width = 20
        ws.column_dimensions["D"].width = 25
        ws.column_dimensions["E"].width = 20

    # ------------------------------------------------------------------
    # Trade logging
    # ------------------------------------------------------------------

    def log_trade(self, trade: Dict):
        """
        Append a single completed trade to the Trades sheet.

        Expected keys in trade dict:
            trade_id, trade_date, entry_time, exit_time, strategy,
            signal_reason, direction, strike, spot_entry, spot_exit,
            entry_premium, exit_premium, sl_amount, tp_amount,
            india_vix, bs_delta, is_expiry_day, orb_high, orb_low,
            vwap_at_entry, quantity, lots, gross_pnl, costs, net_pnl,
            r_value, r_multiple, exit_type
        """
        try:
            wb = load_workbook(str(self.path))
            ws = wb[self.SHEET_TRADES]

            # Find next empty row
            next_row = ws.max_row + 1

            # Compute running equity
            running_equity = self.cfg.capital
            if next_row > 2:
                prev_eq = ws.cell(row=next_row - 1,
                                  column=len(TRADE_COLUMNS)).value
                if prev_eq and isinstance(prev_eq, (int, float)):
                    running_equity = prev_eq
            running_equity += trade.get("net_pnl", 0)

            values = [
                trade.get("trade_id", next_row - 1),
                trade.get("trade_date", ""),
                trade.get("entry_time", ""),
                trade.get("exit_time", ""),
                trade.get("strategy", "sweep"),
                trade.get("signal_reason", ""),
                trade.get("direction", ""),
                trade.get("strike", 0),
                trade.get("spot_entry", 0),
                trade.get("spot_exit", 0),
                trade.get("entry_premium", 0),
                trade.get("exit_premium", 0),
                trade.get("sl_amount", 0),
                trade.get("tp_amount", 0),
                trade.get("india_vix", 0),
                trade.get("bs_delta", 0),
                "Yes" if trade.get("is_expiry_day") else "No",
                trade.get("orb_high", 0),
                trade.get("orb_low", 0),
                trade.get("vwap_at_entry", 0),
                trade.get("quantity", 0),
                trade.get("lots", 0),
                trade.get("gross_pnl", 0),
                trade.get("costs", 0),
                trade.get("net_pnl", 0),
                trade.get("r_value", 0),
                trade.get("r_multiple", 0),
                trade.get("exit_type", ""),
                running_equity,
            ]

            for col_idx, val in enumerate(values, 1):
                cell = ws.cell(row=next_row, column=col_idx, value=val)
                cell.border = _THIN_BORDER
                cell.alignment = Alignment(horizontal="center")

            # Colour the Net P&L cell
            pnl_col = 25  # Net P&L column
            pnl_val = trade.get("net_pnl", 0)
            pnl_cell = ws.cell(row=next_row, column=pnl_col)
            if pnl_val > 0:
                pnl_cell.fill = _GREEN_FILL
            elif pnl_val < 0:
                pnl_cell.fill = _RED_FILL

            # Number formats
            for col in [9, 10, 11, 12, 13, 14, 15, 18, 19, 20]:
                ws.cell(row=next_row, column=col).number_format = _RUPEE_FMT_2
            for col in [23, 24, 25, 29]:
                ws.cell(row=next_row, column=col).number_format = _RUPEE_FMT
            ws.cell(row=next_row, column=26).number_format = _RUPEE_FMT_2
            ws.cell(row=next_row, column=27).number_format = '0.00'

            wb.save(str(self.path))
            wb.close()
            logger.info(f"Trade #{trade.get('trade_id')} logged to Excel")

        except Exception as e:
            logger.error(f"Excel trade log failed: {e}", exc_info=True)

    # ------------------------------------------------------------------
    # Summary generation
    # ------------------------------------------------------------------

    def rebuild_summaries(self):
        """
        Read all trades from the Trades sheet and regenerate:
          - Daily Summary
          - Weekly Summary
          - Monthly Summary
          - Dashboard
        Call this at EOD.
        """
        try:
            wb = load_workbook(str(self.path))
            trades = self._read_all_trades(wb)
            if not trades:
                wb.close()
                return

            self._build_daily_summary(wb, trades)
            self._build_weekly_summary(wb, trades)
            self._build_monthly_summary(wb, trades)
            self._build_dashboard(wb, trades)

            wb.save(str(self.path))
            wb.close()
            logger.info("Excel summaries rebuilt")
        except Exception as e:
            logger.error(f"Summary rebuild failed: {e}", exc_info=True)

    def _read_all_trades(self, wb) -> List[Dict]:
        """Read all trade rows from the Trades sheet."""
        ws = wb[self.SHEET_TRADES]
        trades = []
        headers = [cell.value for cell in ws[1]]

        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] is None:
                continue
            trade = {}
            for idx, header in enumerate(headers):
                if idx < len(row):
                    trade[header] = row[idx]
            trades.append(trade)
        return trades

    def _build_daily_summary(self, wb, trades: List[Dict]):
        """Aggregate trades by date."""
        ws = wb[self.SHEET_DAILY]
        self._clear_data_rows(ws)

        # Group by date
        by_date: Dict[str, List[Dict]] = {}
        for t in trades:
            d = str(t.get("Date", ""))
            if d:
                by_date.setdefault(d, []).append(t)

        running_equity = self.cfg.capital
        row = 2
        for d in sorted(by_date.keys()):
            day_trades = by_date[d]
            n = len(day_trades)
            wins = sum(1 for t in day_trades if (t.get("Net P&L") or 0) > 0)
            losses = n - wins
            win_rate = wins / n if n > 0 else 0
            gross = sum(t.get("Gross P&L") or 0 for t in day_trades)
            net = sum(t.get("Net P&L") or 0 for t in day_trades)
            running_equity += net

            # Max intraday drawdown
            cumulative = 0.0
            peak = 0.0
            max_dd = 0.0
            for t in day_trades:
                cumulative += (t.get("Net P&L") or 0)
                if cumulative > peak:
                    peak = cumulative
                dd = peak - cumulative
                if dd > max_dd:
                    max_dd = dd

            values = [d, n, wins, losses, win_rate, gross, net, max_dd,
                      running_equity]
            for col_idx, val in enumerate(values, 1):
                cell = ws.cell(row=row, column=col_idx, value=val)
                cell.border = _THIN_BORDER
                cell.alignment = Alignment(horizontal="center")

            # Formats
            ws.cell(row=row, column=5).number_format = _PCT_FMT
            for col in [6, 7, 8, 9]:
                ws.cell(row=row, column=col).number_format = _RUPEE_FMT

            # Colour net P&L
            pnl_cell = ws.cell(row=row, column=7)
            if net > 0:
                pnl_cell.fill = _GREEN_FILL
            elif net < 0:
                pnl_cell.fill = _RED_FILL

            row += 1

    def _build_weekly_summary(self, wb, trades: List[Dict]):
        """Aggregate trades by ISO week."""
        ws = wb[self.SHEET_WEEKLY]
        self._clear_data_rows(ws)

        by_week: Dict[str, List[Dict]] = {}
        for t in trades:
            d = str(t.get("Date", ""))
            if not d:
                continue
            try:
                dt = datetime.strptime(d, "%Y-%m-%d").date()
            except ValueError:
                continue
            week_start = dt - timedelta(days=dt.weekday())
            key = week_start.isoformat()
            by_week.setdefault(key, []).append(t)

        running_equity = self.cfg.capital
        row = 2
        for ws_key in sorted(by_week.keys()):
            week_trades = by_week[ws_key]
            week_start = datetime.strptime(ws_key, "%Y-%m-%d").date()
            week_end = week_start + timedelta(days=4)
            n = len(week_trades)
            wins = sum(1 for t in week_trades if (t.get("Net P&L") or 0) > 0)
            losses = n - wins
            win_rate = wins / n if n > 0 else 0
            net = sum(t.get("Net P&L") or 0 for t in week_trades)
            running_equity += net

            r_multiples = [
                t.get("R-Multiple") or 0 for t in week_trades
                if t.get("R-Multiple") is not None
            ]
            avg_r = np.mean(r_multiples) if r_multiples else 0

            values = [
                ws_key, week_end.isoformat(), n, wins, losses,
                win_rate, net, round(avg_r, 2), running_equity,
            ]
            for col_idx, val in enumerate(values, 1):
                cell = ws.cell(row=row, column=col_idx, value=val)
                cell.border = _THIN_BORDER
                cell.alignment = Alignment(horizontal="center")

            ws.cell(row=row, column=6).number_format = _PCT_FMT
            for col in [7, 9]:
                ws.cell(row=row, column=col).number_format = _RUPEE_FMT
            ws.cell(row=row, column=8).number_format = '0.00'

            pnl_cell = ws.cell(row=row, column=7)
            if net > 0:
                pnl_cell.fill = _GREEN_FILL
            elif net < 0:
                pnl_cell.fill = _RED_FILL

            row += 1

    def _build_monthly_summary(self, wb, trades: List[Dict]):
        """Aggregate trades by month."""
        ws = wb[self.SHEET_MONTHLY]
        self._clear_data_rows(ws)

        by_month: Dict[str, List[Dict]] = {}
        for t in trades:
            d = str(t.get("Date", ""))
            if not d:
                continue
            month_key = d[:7]  # "YYYY-MM"
            by_month.setdefault(month_key, []).append(t)

        running_equity = self.cfg.capital
        row = 2
        for mk in sorted(by_month.keys()):
            month_trades = by_month[mk]
            n = len(month_trades)
            wins = sum(1 for t in month_trades if (t.get("Net P&L") or 0) > 0)
            losses = n - wins
            win_rate = wins / n if n > 0 else 0
            pnls = [t.get("Net P&L") or 0 for t in month_trades]
            net = sum(pnls)
            running_equity += net

            gross_profit = sum(p for p in pnls if p > 0)
            gross_loss = abs(sum(p for p in pnls if p < 0))
            pf = gross_profit / gross_loss if gross_loss > 0 else (
                float("inf") if gross_profit > 0 else 0
            )

            # Max drawdown within the month
            cum = 0.0
            peak = 0.0
            max_dd = 0.0
            for p in pnls:
                cum += p
                if cum > peak:
                    peak = cum
                dd = peak - cum
                if dd > max_dd:
                    max_dd = dd

            values = [
                mk, n, wins, losses, win_rate, net,
                round(pf, 3), max_dd, running_equity,
            ]
            for col_idx, val in enumerate(values, 1):
                cell = ws.cell(row=row, column=col_idx, value=val)
                cell.border = _THIN_BORDER
                cell.alignment = Alignment(horizontal="center")

            ws.cell(row=row, column=5).number_format = _PCT_FMT
            for col in [6, 8, 9]:
                ws.cell(row=row, column=col).number_format = _RUPEE_FMT
            ws.cell(row=row, column=7).number_format = '0.000'

            pnl_cell = ws.cell(row=row, column=6)
            if net > 0:
                pnl_cell.fill = _GREEN_FILL
            elif net < 0:
                pnl_cell.fill = _RED_FILL

            row += 1

    def _build_dashboard(self, wb, trades: List[Dict]):
        """Build key metrics dashboard."""
        ws = wb[self.SHEET_DASHBOARD]
        # Clear existing content
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row,
                                max_col=6):
            for cell in row:
                cell.value = None

        pnls = np.array([t.get("Net P&L") or 0 for t in trades])
        n = len(pnls)
        wins_mask = pnls > 0
        n_wins = int(wins_mask.sum())
        n_losses = n - n_wins

        net_total = float(pnls.sum())
        gross_profit = float(pnls[wins_mask].sum()) if n_wins > 0 else 0
        gross_loss = float(abs(pnls[~wins_mask].sum())) if n_losses > 0 else 0
        win_rate = n_wins / n if n > 0 else 0
        pf = gross_profit / gross_loss if gross_loss > 0 else float("inf")
        avg_win = float(pnls[wins_mask].mean()) if n_wins > 0 else 0
        avg_loss = float(pnls[~wins_mask].mean()) if n_losses > 0 else 0
        expectancy = (win_rate * avg_win + (1 - win_rate) * avg_loss) if n > 0 else 0

        # Equity curve
        equity = self.cfg.capital + np.cumsum(pnls)
        peak_eq = np.maximum.accumulate(equity)
        dd_pct = (equity - peak_eq) / peak_eq * 100
        max_dd_pct = float(dd_pct.min()) if len(dd_pct) > 0 else 0
        max_dd_abs = float((peak_eq - equity).max()) if len(equity) > 0 else 0

        # R-multiples
        r_mults = [
            t.get("R-Multiple") or 0 for t in trades
            if t.get("R-Multiple") is not None
        ]
        avg_r = float(np.mean(r_mults)) if r_mults else 0

        # Streaks
        max_win_streak = 0
        max_loss_streak = 0
        cur_win = 0
        cur_loss = 0
        for p in pnls:
            if p > 0:
                cur_win += 1
                cur_loss = 0
                if cur_win > max_win_streak:
                    max_win_streak = cur_win
            else:
                cur_loss += 1
                cur_win = 0
                if cur_loss > max_loss_streak:
                    max_loss_streak = cur_loss

        # Strategy breakdown
        sweep_pnl = sum(
            t.get("Net P&L") or 0 for t in trades
            if (t.get("Strategy") or "").lower() in ("sweep", "orb_sweep")
        )
        straddle_pnl = sum(
            t.get("Net P&L") or 0 for t in trades
            if (t.get("Strategy") or "").lower() in ("straddle", "short_straddle")
        )

        # Write dashboard
        r = 1
        ws.cell(row=r, column=1, value="ORB Sweep + Straddle Dashboard").font = _TITLE_FONT
        r += 2

        metrics_left = [
            ("Total Trades", n),
            ("Wins / Losses", f"{n_wins} / {n_losses}"),
            ("Win Rate", f"{win_rate*100:.1f}%"),
            ("Profit Factor", f"{pf:.3f}" if pf != float("inf") else "Inf"),
            ("Expectancy (per trade)", f"Rs {expectancy:+,.0f}"),
            ("", ""),
            ("Net P&L (Total)", f"Rs {net_total:+,.0f}"),
            ("Gross Profit", f"Rs {gross_profit:+,.0f}"),
            ("Gross Loss", f"Rs {-gross_loss:+,.0f}"),
            ("Avg Win", f"Rs {avg_win:+,.0f}"),
            ("Avg Loss", f"Rs {avg_loss:+,.0f}"),
            ("", ""),
            ("Max Drawdown %", f"{max_dd_pct:.2f}%"),
            ("Max Drawdown Rs", f"Rs {max_dd_abs:,.0f}"),
            ("Current Equity", f"Rs {float(equity[-1]):,.0f}" if len(equity) > 0 else "N/A"),
            ("Return on Capital", f"{net_total/self.cfg.capital*100:.2f}%"),
        ]

        metrics_right = [
            ("Avg R-Multiple", f"{avg_r:.2f}R"),
            ("Max Win Streak", max_win_streak),
            ("Max Loss Streak", max_loss_streak),
            ("", ""),
            ("Strategy Breakdown", ""),
            ("  Sweep Net P&L", f"Rs {sweep_pnl:+,.0f}"),
            ("  Straddle Net P&L", f"Rs {straddle_pnl:+,.0f}"),
            ("", ""),
            ("Capital", f"Rs {self.cfg.capital:,.0f}"),
            ("Risk per Trade", f"{self.cfg.risk_per_trade_pct*100:.1f}%"),
            ("Daily DD Limit", f"{self.cfg.daily_dd_limit_pct*100:.1f}%"),
            ("Weekly DD Limit", f"{self.cfg.weekly_dd_limit_pct*100:.1f}%"),
        ]

        for i, (label, val) in enumerate(metrics_left):
            row_n = r + i
            lbl_cell = ws.cell(row=row_n, column=1, value=label)
            val_cell = ws.cell(row=row_n, column=2, value=val)
            if label:
                lbl_cell.font = _BOLD_FONT
            val_cell.alignment = Alignment(horizontal="right")

        for i, (label, val) in enumerate(metrics_right):
            row_n = r + i
            lbl_cell = ws.cell(row=row_n, column=4, value=label)
            val_cell = ws.cell(row=row_n, column=5, value=val)
            if label:
                lbl_cell.font = _BOLD_FONT
            val_cell.alignment = Alignment(horizontal="right")

    # ------------------------------------------------------------------
    # Utility
    # ------------------------------------------------------------------

    @staticmethod
    def _clear_data_rows(ws):
        """Remove all rows below the header."""
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row)

    def get_trade_count(self) -> int:
        """Return number of trades currently in the journal."""
        try:
            wb = load_workbook(str(self.path), read_only=True)
            ws = wb[self.SHEET_TRADES]
            count = ws.max_row - 1  # subtract header
            wb.close()
            return max(0, count)
        except Exception:
            return 0


# ---------------------------------------------------------------------------
# Helper: compute BS delta for logging
# ---------------------------------------------------------------------------

def bs_delta(spot: float, strike: int, vix: float,
             minutes_to_expiry: float, opt_type: str,
             rfr: float = 0.065) -> float:
    """
    Black-Scholes delta for an option.
    Returns delta in [0, 1] for calls, [-1, 0] for puts.
    If minutes_to_expiry is 0 or vix is 0, returns intrinsic delta.
    """
    try:
        from scipy.stats import norm
    except ImportError:
        # Rough approximation if scipy not available
        if opt_type.upper() in ("CE", "CALL"):
            return 0.5 if spot >= strike else 0.3
        else:
            return -0.5 if spot <= strike else -0.3

    if minutes_to_expiry <= 0 or vix <= 0:
        if opt_type.upper() in ("CE", "CALL"):
            return 1.0 if spot > strike else 0.0
        else:
            return -1.0 if spot < strike else 0.0

    T = max(minutes_to_expiry / (365 * 375), 1e-8)
    sigma = max(vix / 100, 0.01)
    S, K = float(spot), float(strike)

    d1 = (math.log(S / K) + (rfr + 0.5 * sigma ** 2) * T) / (sigma * math.sqrt(T))

    if opt_type.upper() in ("CE", "CALL"):
        return float(norm.cdf(d1))
    else:
        return float(norm.cdf(d1) - 1)
