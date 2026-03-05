"""Compare backtest vs live trade data to identify red flags."""
import pandas as pd
import json

# Load backtest trades
df = pd.read_csv("reports/3year_backtest/all_trades.csv")
print("Backtest columns:", list(df.columns))
print(f"Total backtest trades: {len(df)}")
print()

# Per-symbol analysis for traded symbols
for sym in ["US2000", "SWI20"]:
    sub = df[df["symbol"] == sym]
    print(f"=== BACKTEST: {sym} ({len(sub)} trades) ===")
    
    # Slippage / cost columns
    for col in ["slippage", "slippage_pts", "spread", "entry_spread", "exit_spread", 
                 "cost", "total_cost", "spread_cost", "commission"]:
        if col in df.columns:
            print(f"  {col}: mean={sub[col].mean():.3f}, std={sub[col].std():.3f}, max={sub[col].max():.3f}")
    
    # RR distribution
    if "actual_rr" in df.columns:
        print(f"  actual_rr: mean={sub['actual_rr'].mean():.3f}, median={sub['actual_rr'].median():.3f}")
    
    # P&L
    if "pnl" in df.columns:
        wins = sub[sub["pnl"] > 0]
        losses = sub[sub["pnl"] < 0]
        bes = sub[sub["pnl"] == 0]
        print(f"  avg_win=${wins['pnl'].mean():.2f}, avg_loss=${losses['pnl'].mean():.2f}")
        print(f"  win_rate={len(wins)/len(sub)*100:.1f}%, BE={len(bes)/len(sub)*100:.1f}%")
    
    # Time in trade
    for col in ["duration", "time_in_trade", "bars_held", "duration_mins"]:
        if col in df.columns:
            vals = pd.to_numeric(sub[col], errors="coerce")
            print(f"  {col}: mean={vals.mean():.1f}, median={vals.median():.1f}")
    
    # Exit reasons
    if "exit_reason" in df.columns:
        print(f"  exit_reasons: {sub['exit_reason'].value_counts().to_dict()}")
    
    # SL hit analysis  
    if "exit_reason" in df.columns and "pnl" in df.columns:
        sl_hits = sub[sub["exit_reason"] == "SL_HIT"]
        if len(sl_hits) > 0:
            print(f"  SL_HIT: {len(sl_hits)} trades, avg_loss=${sl_hits['pnl'].mean():.2f}")
    
    # Quick reversal analysis (stopped out fast)
    if "bars_held" in df.columns:
        fast = sub[sub["bars_held"] <= 5]
        if len(fast) > 0:
            fast_loss = fast[fast["pnl"] < 0]
            print(f"  Fast stops (<=5 bars): {len(fast)}/{len(sub)} = {len(fast)/len(sub)*100:.1f}%")
            if len(fast_loss) > 0:
                print(f"    of which losses: {len(fast_loss)} avg=${fast_loss['pnl'].mean():.2f}")
    
    # ORB range if available
    if "orb_range" in df.columns:
        print(f"  ORB range: mean={sub['orb_range'].mean():.2f}, std={sub['orb_range'].std():.2f}")
    
    print()

# Overall portfolio stats
print("=== BACKTEST: OVERALL PORTFOLIO ===")
metrics = json.load(open("reports/3year_backtest/metrics_3year.json"))
ov = metrics["overall"]
print(f"  Win rate: {ov['win_rate']}%")
print(f"  Profit factor: {ov['profit_factor']}")
print(f"  Avg win: ${ov['avg_win']:.2f}, Avg loss: ${ov['avg_loss']:.2f}")
print(f"  Expectancy: ${ov['expectancy']:.2f}/trade")
print(f"  Max DD: {ov['max_dd_pct']}%")
print(f"  Trades/week: {ov['trades_per_week']}")
er = metrics["exit_reasons"]
for reason, stats in er.items():
    print(f"  {reason}: {stats['count']} ({stats['pct']}%), avg_pnl=${stats['avg_pnl']:.2f}")

print()
print("=" * 60)
print("=== LIVE TRADES ===")
print("=" * 60)

# Live trades
import openpyxl
wb = openpyxl.load_workbook("live/reports/trade_journal.xlsx")
ws = wb["Trades"]
headers = [c.value for c in ws[1]]

live_trades = []
for row in ws.iter_rows(min_row=2, values_only=True):
    if any(v is not None for v in row):
        live_trades.append(dict(zip(headers, row)))

for t in live_trades:
    tid = t["Trade ID"]
    sym = t["Symbol"]
    print(f"\nTrade #{tid} - {sym} {t['Direction']}")
    print(f"  Entry: planned={t['Planned Entry']}, actual={t['Actual Entry']}, slippage={t['Slippage Pts']} pts (${t.get('Slippage $',0) or 0:.2f})")
    print(f"  SL={t['Stop Loss']}, TP={t['TP']}, R={t['R (Risk Pts)']} pts")
    print(f"  ORB Range: {t.get('ORB Range',0)}")
    print(f"  Exit: {t['Exit Price']} ({t['Exit Reason']})")
    print(f"  P&L: ${t.get('P&L',0) or 0:.2f}, RR: {t.get('Actual RR',0) or 0:.2f}")
    print(f"  MFE(R): {t.get('MFE (R)','?')}, MAE(R): {t.get('MAE (R)','?')}")
    print(f"  Time in trade: {t.get('Time in Trade','?')}")
    print(f"  Exec delay: {t.get('Exec Delay (s)',0)}s, Spread: {t.get('Spread Entry',0)}")

print()
print("=" * 60)
print("=== RED FLAG ANALYSIS ===")
print("=" * 60)

# Calculations
live_wins = sum(1 for t in live_trades if (t.get("P&L",0) or 0) > 0)
live_losses = sum(1 for t in live_trades if (t.get("P&L",0) or 0) < 0)
live_bes = sum(1 for t in live_trades if (t.get("P&L",0) or 0) == 0)
live_total = len(live_trades)
live_wr = live_wins / live_total * 100 if live_total > 0 else 0

live_pnl = sum(t.get("P&L",0) or 0 for t in live_trades)
live_avg_pnl = live_pnl / live_total if live_total > 0 else 0

# Slippage analysis
slippages = [(t["Symbol"], t.get("Slippage Pts",0) or 0, t.get("Slippage $",0) or 0) for t in live_trades]
avg_slip_dollar = sum(s[2] for s in slippages) / len(slippages) if slippages else 0

print(f"\n1. SAMPLE SIZE FLAG")
print(f"   Live trades: {live_total} (need 30+ for statistical significance)")
print(f"   Confidence: VERY LOW - cannot draw conclusions from {live_total} trades")

print(f"\n2. WIN RATE")
print(f"   Live: {live_wr:.1f}% ({live_wins}W / {live_losses}L / {live_bes}BE)")
print(f"   Backtest overall: {ov['win_rate']}%")
bt_us2000_wr = metrics['per_index']['US2000']['win_rate']
bt_swi20_wr = metrics['per_index']['SWI20']['win_rate']
print(f"   Backtest US2000: {bt_us2000_wr}%")
print(f"   Backtest SWI20: {bt_swi20_wr}%")

print(f"\n3. SLIPPAGE")
for sym, spts, sdol in slippages:
    direction = "adverse" if spts > 0 else "favorable"
    print(f"   Trade {sym}: {spts} pts (${sdol:.2f}) - {direction}")
print(f"   Average slippage: ${avg_slip_dollar:.2f}")
print(f"   Backtest assumed cost per trade: ${ov['total_costs']/ov['total_trades']:.2f}")

print(f"\n4. EXECUTION DELAYS")
for t in live_trades:
    delay = t.get("Exec Delay (s)", 0) or 0
    flag = "CRITICAL" if delay > 60 else "OK" if delay < 5 else "WARNING"
    print(f"   Trade #{t['Trade ID']} ({t['Symbol']}): {delay}s [{flag}]")

print(f"\n5. RR ACHIEVED vs PLANNED")
for t in live_trades:
    planned = t.get("Planned RR", 0) or 0
    actual = t.get("Actual RR", 0) or 0
    pnl_val = t.get("P&L", 0) or 0
    print(f"   Trade #{t['Trade ID']} ({t['Symbol']}): planned={planned}, actual={actual}, P&L=${pnl_val:.2f}")

print(f"\n6. TIME IN TRADE")
for t in live_trades:
    print(f"   Trade #{t['Trade ID']} ({t['Symbol']}): {t.get('Time in Trade','?')}")
bt_bars = df[df["symbol"].isin(["US2000","SWI20"])]
if "bars_held" in df.columns:
    avg_bars = bt_bars["bars_held"].mean()
    print(f"   Backtest avg bars held: {avg_bars:.0f}")

print(f"\n7. MFE/MAE ANALYSIS (Edge Capture)")
for t in live_trades:
    mfe_r = t.get("MFE (R)", None)
    mae_r = t.get("MAE (R)", None)
    actual_rr = t.get("Actual RR", 0) or 0
    if mfe_r is not None and mfe_r != 0:
        capture = actual_rr / mfe_r * 100 if mfe_r else 0
        print(f"   Trade #{t['Trade ID']} ({t['Symbol']}): MFE={mfe_r}R, MAE={mae_r}R, Captured={capture:.0f}% of max move")
    else:
        print(f"   Trade #{t['Trade ID']} ({t['Symbol']}): MFE={mfe_r}R, MAE={mae_r}R")

print(f"\n8. SWI20 CONCERN")
bt_swi20 = metrics['per_index']['SWI20']
print(f"   Backtest PF: {bt_swi20['profit_factor']} (< 1.0 = LOSING)")
print(f"   Backtest expectancy: ${bt_swi20['expectancy']:.2f}/trade")
print(f"   Backtest net P&L: ${bt_swi20['net_pnl']:.2f}")
print(f"   Live result: 1 trade, -$46.27 (SL hit)")
print(f"   VERDICT: SWI20 is a net loser in backtest - should reconsider trading it")

print(f"\n9. RISK PER TRADE")
for t in live_trades:
    risk = t.get("Risk Amount", 0) or 0
    bal_at_trade = (t.get("Balance After", 0) or 0) - (t.get("P&L", 0) or 0)
    risk_pct = risk / bal_at_trade * 100 if bal_at_trade else 0
    r_pts = t.get("R (Risk Pts)", 0) or 0
    orb = t.get("ORB Range", 0) or 0
    r_orb_ratio = r_pts / orb if orb else 0
    print(f"   Trade #{t['Trade ID']} ({t['Symbol']}): risk=${risk:.2f} ({risk_pct:.2f}%), R={r_pts:.1f}pts, ORB={orb}, R/ORB={r_orb_ratio:.2f}x")

print(f"\n10. NET P&L")
print(f"   Live: ${live_pnl:.2f} on {live_total} trades")
print(f"   Backtest expectancy: ${ov['expectancy']:.2f}/trade")
print(f"   Expected for {live_total} trades: ${ov['expectancy']*live_total:.2f}")
print(f"   Actual vs Expected: ${live_pnl - ov['expectancy']*live_total:.2f}")
