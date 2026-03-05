import openpyxl
wb = openpyxl.load_workbook('live/reports/trade_journal.xlsx')
ws = wb.active
headers = [cell.value for cell in ws[1]]
print(f"Total trades: {ws.max_row - 1}\n")
for row in ws.iter_rows(min_row=2, values_only=True):
    d = dict(zip(headers, row))
    print(f"Trade #{d.get('Trade ID')} | {d.get('Date')} {d.get('Time')} | {d.get('Symbol')} {d.get('Direction')}")
    print(f"  Planned Entry: {d.get('Planned Entry')}  Actual Entry: {d.get('Actual Entry')}  SL: {d.get('Stop Loss')}  TP: {d.get('TP')}")
    print(f"  R (risk pts): {d.get('R (Risk Pts)')}  Lots: {d.get('Lot Size')}  Risk $: {d.get('Risk Amount')}")
    print(f"  Exit: {d.get('Exit Price')}  Reason: {d.get('Exit Reason')}  PnL: ${d.get('P&L')}  Actual RR: {d.get('Actual RR')}")
    print(f"  Slippage: {d.get('Slippage Pts')} pts  Spread: {d.get('Spread Entry')}")
    print(f"  MFE: {d.get('MFE (pts)')} pts  MAE: {d.get('MAE (pts)')} pts")
    print(f"  Balance After: {d.get('Balance After')}  Session: {d.get('Session')}  Weekday: {d.get('Weekday')}")
    print()

# Also read edge guard
import json
with open('live/logs/edge_guard.json') as f:
    eg = json.load(f)
trades = eg.get('trades', [])
real_trades = [t for t in trades if t.get('symbol')]
print(f"\n--- Edge Guard: {len(real_trades)} real trades ---")
for t in real_trades:
    print(f"  {t.get('date')} | {t.get('symbol'):8s} | {t.get('direction'):7s} | PnL=${t.get('pnl', 0):+.2f} | RR={t.get('actual_rr', 0):+.2f} | Bal={t.get('balance_after', 0):.2f}")
